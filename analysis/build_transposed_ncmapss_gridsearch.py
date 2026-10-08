

import argparse
import glob
import os
import re
import statistics as stats

import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

DEFAULT_DATA_DIR = r"./external_data/NCMAPSS_logs_gs"
OUT_PATH = os.path.join(os.path.dirname(os.path.abspath(__file__)), "out", "NCMAPSS_gridsearch_transposed.xlsx")


MODEL_ORDER = [
    "LBGN_RUL", "PDMN", "NDC_PDMN",
    "DAGCN_RUL", "EviAdaptRUL", "DAST_RUL", "CADA_RUL", "TACDA_RUL",
    "OCS_DANN", "MDAN_RUL", "CCDG_RUL",
]
OURS_KEYS = {"LBGN_RUL", "PDMN", "NDC_PDMN"}

FONT_NAME = "Arial"
THIN_GRAY = Side(style="thin", color="BFBFBF")
BORDER_ALL = Border(left=THIN_GRAY, right=THIN_GRAY, top=THIN_GRAY, bottom=THIN_GRAY)

TITLE_FONT = Font(name=FONT_NAME, size=8, italic=True, color="595959")
TITLE_ALIGN = Alignment(horizontal="left", vertical="center", wrap_text=True)

HEADER_FONT = Font(name=FONT_NAME, size=9, bold=True, color="FFFFFF")
HEADER_FILL = PatternFill(fill_type="solid", fgColor="1F3864")
HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)

MODEL_GROUP_FONT = Font(name=FONT_NAME, size=9, bold=True, color="1F3864")
MODEL_GROUP_FILL = PatternFill(fill_type="solid", fgColor="D9E1F2")

ROW_LABEL_FONT = Font(name=FONT_NAME, size=9, color="000000")
ROW_LABEL_ALIGN = Alignment(horizontal="left", vertical="center")

DATA_FONT = Font(name=FONT_NAME, size=9, color="000000")
DATA_ALIGN = Alignment(horizontal="center", vertical="center")

DASH_FONT = Font(name=FONT_NAME, size=9, italic=True, color="808080")
DASH_FILL = PatternFill(fill_type="solid", fgColor="F2F2F2")

MEAN_FONT = Font(name=FONT_NAME, size=9, bold=True, color="000000")
BEST_FILL = PatternFill(fill_type="solid", fgColor="C6EFCE")


def discover_pairs(data_dir):

    pairs = []
    for fpath in glob.glob(os.path.join(data_dir, "*experimental_logs.csv")):
        base = os.path.basename(fpath)
        pair = base[: -len("experimental_logs.csv")]
        src, _, tgt = pair.partition("_")
        pairs.append((src, tgt))
    return sorted(set(pairs))


def load_all(data_dir, pairs):
    frames = []
    for src, tgt in pairs:
        fpath = os.path.join(data_dir, f"{src}_{tgt}experimental_logs.csv")
        df = pd.read_csv(fpath)
        df["pair"] = f"{src}_{tgt}"
        frames.append(df)
    full = pd.concat(frames, ignore_index=True)
    full = full[full["info"].fillna("").str.contains("_gridsearch")].copy()


    full["seed"] = full["info"].str.extract(r"seed(\d+)")
    return full


def row_key(row):

    return (row["model"], row["LR"], row["hidden_dim"])


def build_sheet(wb, sheet_name, title_text, full, pairs, value_col):
    ws = wb.create_sheet(sheet_name)
    pair_labels = [f"{s}->{t}" for s, t in pairs]
    last_col = 2 + len(pairs) + 1

    ws.column_dimensions["A"].width = 28
    for c in range(2, last_col + 1):
        ws.column_dimensions[get_column_letter(c)].width = 13

    row = 1
    ws.cell(row=row, column=1, value=title_text)
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=last_col)
    ws.cell(row=row, column=1).font = TITLE_FONT
    ws.cell(row=row, column=1).alignment = TITLE_ALIGN
    ws.row_dimensions[row].height = 30
    row += 1

    ws.cell(row=row, column=1, value="Model / LR / hidden_dim")
    for i, label in enumerate(pair_labels):
        ws.cell(row=row, column=2 + i, value=label)
    ws.cell(row=row, column=last_col, value="Overall\nmean")
    for c in range(1, last_col + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = HEADER_ALIGN
    ws.row_dimensions[row].height = 26
    row += 1

    for model in MODEL_ORDER:
        sub = full[full["model"] == model]
        if sub.empty:
            continue


        grid_points = sorted(
            {(r["LR"], r["hidden_dim"]) for _, r in sub.iterrows()},
            key=lambda x: (-x[0], x[1]),
        )

        group_row = row
        label = "LBGN (ours)" if model == "LBGN_RUL" else model
        ws.cell(row=group_row, column=1, value=label)
        ws.cell(row=group_row, column=1).font = MODEL_GROUP_FONT
        ws.cell(row=group_row, column=1).fill = MODEL_GROUP_FILL
        ws.merge_cells(start_row=group_row, start_column=1, end_row=group_row, end_column=last_col)
        row += 1

        first_data_row = row
        for lr, hd in grid_points:
            hd_suffix = f"  hd={int(hd)}" if model == "LBGN_RUL" else ""
            ws.cell(row=row, column=1, value=f"  lr={lr:g}{hd_suffix}")
            ws.cell(row=row, column=1).font = ROW_LABEL_FONT
            ws.cell(row=row, column=1).alignment = ROW_LABEL_ALIGN

            vals_for_mean = []
            for i, (src, tgt) in enumerate(pairs):
                pair = f"{src}_{tgt}"
                cell_rows = sub[(sub["pair"] == pair) & (sub["LR"] == lr) & (sub["hidden_dim"] == hd)]
                cell = ws.cell(row=row, column=2 + i)
                cell.alignment = DATA_ALIGN
                if cell_rows.empty:
                    cell.value = "-"
                    cell.font = DASH_FONT
                    cell.fill = DASH_FILL
                else:
                    v = float(cell_rows.iloc[-1][value_col])
                    cell.value = round(v, 3) if value_col == "best_last_RMSE" else round(v, 4)
                    cell.font = DATA_FONT
                    vals_for_mean.append(v)

            mcell = ws.cell(row=row, column=last_col)
            mcell.alignment = DATA_ALIGN
            if vals_for_mean:
                mcell.value = round(stats.mean(vals_for_mean), 4)
                mcell.font = MEAN_FONT
            else:
                mcell.value = "-"
                mcell.font = DASH_FONT
            row += 1
        last_data_row = row - 1


        best_row, best_val = None, None
        for r in range(first_data_row, last_data_row + 1):
            v = ws.cell(row=r, column=last_col).value
            if v in (None, "-"):
                continue
            if best_val is None or v < best_val:
                best_val, best_row = v, r
        if best_row is not None:
            for c in range(1, last_col + 1):
                ws.cell(row=best_row, column=c).fill = BEST_FILL

        row += 1
        ws.row_dimensions[row - 1].height = 6

    for r in range(1, row):
        for c in range(1, last_col + 1):
            if ws.cell(row=r, column=c).value is not None:
                ws.cell(row=r, column=c).border = BORDER_ALL

    ws.freeze_panes = "B3"
    return ws


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--data_dir", default=DEFAULT_DATA_DIR)
    ap.add_argument("--out", default=OUT_PATH)
    args = ap.parse_args()

    pairs = discover_pairs(args.data_dir)
    print(f"Pairs found ({len(pairs)}): {[f'{s}->{t}' for s, t in pairs]}")
    missing_sources = sorted(({s for s, _ in pairs} | {t for _, t in pairs}) - {s for s, _ in pairs})
    if missing_sources:
        print(f"NOTE: no pair has these as SOURCE yet: {missing_sources} "
              f"(expected: DS01 rerun in progress)")

    full = load_all(args.data_dir, pairs)
    print(f"{len(full)} gridsearch rows loaded across {len(pairs)} pair(s), "
          f"{full['model'].nunique()} model(s)")

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    build_sheet(
        wb, "RMSE",
        "N-CMAPSS LR/hidden_dim grid search -- RMSE per (model, LR[, hd]) grid point. "
        "Single seed per cell (see 'info' column in the source CSVs for the exact seed). "
        "'-' = that grid point hasn't finished for this pair yet. Green = lowest overall-mean "
        "row per model -- candidate LR to lock into configs/hparams.py.",
        full, pairs, value_col="best_last_RMSE",
    )
    build_sheet(
        wb, "Score",
        "N-CMAPSS LR/hidden_dim grid search -- PHM08 score per (model, LR[, hd]) grid point "
        "(lower = better, heavy-tailed). Same grid points/highlighting as the RMSE sheet.",
        full, pairs, value_col="score",
    )

    os.makedirs(os.path.dirname(args.out), exist_ok=True)
    wb.save(args.out)
    print(f"\nSaved: {args.out}")


if __name__ == "__main__":
    main()
